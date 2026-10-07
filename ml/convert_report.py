from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
import re

def parse_markdown(file_path):
    """解析markdown文件内容"""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    sections = []
    current_section = None
    
    lines = content.split('\n')
    for line in lines:
        # 一级标题
        if line.startswith('## 一、') or line.startswith('## 二、') or line.startswith('## 三、') or \
           line.startswith('## 四、') or line.startswith('## 五、') or line.startswith('## 六、') or \
           line.startswith('## 七、') or line.startswith('## 八、') or line.startswith('## 九、') or \
           line.startswith('## 十、'):
            if current_section:
                sections.append(current_section)
            current_section = {
                'title': line[3:],
                'content': []
            }
        # 二级标题
        elif line.startswith('### '):
            if current_section:
                current_section['content'].append(('subtitle', line[4:]))
        # 表格
        elif '|' in line and ('---' in line or current_section and current_section['content'] and current_section['content'][-1][0] == 'table'):
            if current_section:
                if not current_section['content'] or current_section['content'][-1][0] != 'table':
                    current_section['content'].append(('table', [line]))
                else:
                    current_section['content'][-1][1].append(line)
        # 列表项
        elif line.startswith('- **'):
            if current_section:
                current_section['content'].append(('list', line[2:]))
        # 代码块
        elif line.startswith('```'):
            if current_section:
                if not current_section['content'] or current_section['content'][-1][0] != 'code':
                    current_section['content'].append(('code_start', ''))
                else:
                    current_section['content'].append(('code_end', ''))
        # 普通段落
        elif line.strip() and not line.startswith('***') and not line.startswith('|'):
            if current_section:
                current_section['content'].append(('paragraph', line.strip()))
    
    if current_section:
        sections.append(current_section)
    
    return sections

def create_report_with_template():
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 设置页面大小为A4
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    
    # ========== 封面页 ==========
    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("人工智能创意大赛项目报告")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.name = "黑体"
    title_run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    # 空行
    for _ in range(3):
        doc.add_paragraph()
    
    # 项目名称
    project_name = doc.add_paragraph()
    project_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    project_name_run = project_name.add_run("项目名称: ")
    project_name_run.font.size = Pt(12)
    project_name_run.font.name = "宋体"
    project_name_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    project_name.add_run("智能校园食堂菜品推荐系统")
    
    # 空行
    for _ in range(2):
        doc.add_paragraph()
    
    # 团队名称
    team_name = doc.add_paragraph()
    team_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    team_name_run = team_name.add_run("团队名称: ")
    team_name_run.font.size = Pt(12)
    team_name_run.font.name = "宋体"
    team_name_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    team_name.add_run("AI美食探索者")
    
    # 空行
    for _ in range(2):
        doc.add_paragraph()
    
    # 成员信息标题
    member_title = doc.add_paragraph()
    member_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    member_title_run = member_title.add_run("成员信息:")
    member_title_run.font.size = Pt(12)
    member_title_run.font.name = "宋体"
    member_title_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 成员表格
    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    
    # 表头
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "序号"
    hdr_cells[1].text = "学号"
    hdr_cells[2].text = "姓名"
    hdr_cells[3].text = "职责"
    
    # 数据
    member_data = [
        ["1", "-", "-", "项目负责人、算法设计"],
        ["2", "-", "-", "前端开发、UI设计"],
        ["3", "-", "-", "后端开发、API设计"],
        ["4", "-", "-", "数据处理、模型训练"]
    ]
    
    for i, data in enumerate(member_data):
        table.rows[i+1].cells[0].text = data[0]
        table.rows[i+1].cells[1].text = data[1]
        table.rows[i+1].cells[2].text = data[2]
        table.rows[i+1].cells[3].text = data[3]
    
    # 设置表格列宽
    table.columns[0].width = Cm(1.5)
    table.columns[1].width = Cm(3)
    table.columns[2].width = Cm(2)
    table.columns[3].width = Cm(8)
    
    # 设置表格字体和对齐
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraphs = cell.paragraphs
            for paragraph in paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = "宋体"
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 空行
    for _ in range(5):
        doc.add_paragraph()
    
    # 完成日期
    date_paragraph = doc.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_paragraph.add_run("完成日期: 2026年4月28日")
    date_run.font.size = Pt(12)
    date_run.font.name = "宋体"
    date_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ========== 目录页（新页面） ==========
    doc.add_section(WD_SECTION.NEW_PAGE)
    
    # 目录标题
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    toc_title_run = toc_title.add_run("参考目录")
    toc_title_run.font.size = Pt(12)
    toc_title_run.font.name = "宋体"
    toc_title_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 目录内容
    toc_items = [
        "一、项目概述",
        "二、项目背景与问题分析",
        "三、项目创意设计与人工智能应用",
        "四、系统总体设计与实现思路",
        "五、项目创新点",
        "六、可行性分析",
        "七、应用价值与预期效果",
        "八、团队分工与实施计划",
        "九、总结与未来展望",
        "十、参考文献"
    ]
    
    for item in toc_items:
        toc_item = doc.add_paragraph()
        toc_item.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        toc_item_run = toc_item.add_run(item)
        toc_item_run.font.size = Pt(12)
        toc_item_run.font.name = "宋体"
        toc_item_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 解析markdown内容
    sections = parse_markdown('项目报告.md')
    
    # 生成正文内容
    for sec in sections:
        doc.add_section(WD_SECTION.NEW_PAGE)
        
        # 章节标题
        chapter = doc.add_paragraph()
        chapter_run = chapter.add_run(sec['title'])
        chapter_run.font.size = Pt(16)
        chapter_run.font.bold = True
        chapter_run.font.name = "黑体"
        chapter_run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        
        # 章节内容
        for content_type, content in sec['content']:
            if content_type == 'subtitle':
                doc.add_paragraph()
                p = doc.add_paragraph()
                p_run = p.add_run(content)
                p_run.font.size = Pt(12)
                p_run.font.bold = True
                p_run.font.name = "宋体"
                p_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            elif content_type == 'paragraph':
                doc.add_paragraph()
                p = doc.add_paragraph()
                p_run = p.add_run(content)
                p_run.font.size = Pt(12)
                p_run.font.name = "宋体"
                p_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                p.first_line_indent = Cm(0.74)
            elif content_type == 'list':
                p = doc.add_paragraph()
                p_run = p.add_run("• " + content)
                p_run.font.size = Pt(12)
                p_run.font.name = "宋体"
                p_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                p.first_line_indent = Cm(0.74)
            elif content_type == 'table':
                doc.add_paragraph()
                # 解析表格
                rows = content
                if len(rows) >= 2:
                    # 确定列数
                    cols = len(rows[0].split('|')) - 2
                    table = doc.add_table(rows=len(rows)-1, cols=cols)
                    table.style = "Table Grid"
                    
                    for i, row in enumerate(rows):
                        if i == 1:  # 跳过分隔线
                            continue
                        cells = row.split('|')[1:-1]
                        for j, cell in enumerate(cells):
                            actual_row = i - (1 if i > 1 else 0)
                            table.rows[actual_row].cells[j].text = cell.strip()
                    
                    # 设置表格字体
                    for row in table.rows:
                        for cell in row.cells:
                            paragraphs = cell.paragraphs
                            for paragraph in paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(10)
                                    run.font.name = "宋体"
                                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            elif content_type == 'code_start':
                doc.add_paragraph()
            elif content_type == 'code_end':
                pass
    
    # 保存文档
    doc.save("项目报告_按模板格式.docx")
    print("文档已保存为: 项目报告_按模板格式.docx")

if __name__ == "__main__":
    create_report_with_template()